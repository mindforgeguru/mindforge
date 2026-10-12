"""
Presentation changes reach the rest of the staff room.

The library and the dashboard's "Presentation progress" card are school-wide:
every teacher sees every deck, and one progress card per teacher using it. But
the only event the server sent was `presentation_ready`, and only to the
uploader. A deck uploaded, adopted, edited, taken off a dashboard or deleted by
one teacher left every other teacher's screen as it was until they refreshed.

Decks here are uploaded as .pptx, which is parsed, not generated — no AI call.
Logging a period is deliberately not exercised: it schedules an auto-quiz,
which is one.
"""

import asyncio
import io

import asyncpg
import pytest

from .conftest import DB_DSN, PREFIX, auth
from .test_realtime_delivery import (
    RECV_TIMEOUT, SILENCE_TIMEOUT, _collect, _connect, _mint_access_token,
)

EVENTS = ("presentation_ready", "presentation_updated")


def _pptx(slides: int = 3) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    for i in range(slides):
        s = deck.slides.add_slide(deck.slide_layouts[1])
        s.shapes.title.text = f"Slide {i + 1}"
        s.placeholders[1].text = f"Point {i + 1}"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def _upload(api, token, name):
    r = api.post(
        "/api/presentations/upload-deck",
        headers=auth(token),
        data={"grade": 8, "subject": "Science", "chapter_name": name},
        files={"file": ("deck.pptx", _pptx(), "application/zip")},
    )
    assert r.status_code == 201, r.text
    return r.json()["presentation_id"]


@pytest.fixture(scope="module")
def staff(two_schools):
    """user ids of school A's admin (the colleague who is watching) and school
    B's teacher (the outsider)."""
    async def _ids():
        conn = await asyncpg.connect(DB_DSN)
        try:
            return {
                "colleague": await conn.fetchval(
                    "SELECT id FROM users WHERE username = $1 AND school_id = $2",
                    two_schools["a"]["admin_username"], two_schools["a"]["id"],
                ),
                "outsider": two_schools["b"]["teacher_id"],
            }
        finally:
            await conn.close()

    ids = asyncio.run(_ids())
    assert ids["colleague"] is not None
    return ids


def _heard(listener_id, act, timeout=RECV_TIMEOUT):
    """Run `act()` while `listener_id` is connected; return the presentation
    events that socket received."""
    async def scenario():
        async with _connect(listener_id, _mint_access_token(listener_id)) as ws:
            await asyncio.sleep(0.5)
            act()
            return await _collect(ws, timeout)

    return [e for e in asyncio.run(scenario()) if e.get("event") in EVENTS]


@pytest.fixture(scope="module")
def deck(api, two_schools):
    return _upload(api, two_schools["a"]["teacher_token"], f"{PREFIX} shared deck")


def test_a_colleague_hears_about_a_new_deck(api, two_schools, staff):
    token = two_schools["a"]["teacher_token"]
    got = _heard(staff["colleague"], lambda: _upload(api, token, f"{PREFIX} new deck"))
    assert got, "a new deck in the school library reached nobody but its uploader"


def test_a_colleague_hears_when_a_deck_is_adopted(api, two_schools, staff, deck):
    teacher = auth(two_schools["a"]["teacher_token"])

    def adopt():
        assert api.post(f"/api/presentations/{deck}/adopt", headers=teacher).status_code == 201

    got = _heard(staff["colleague"], adopt)
    assert got and got[0]["presentation_id"] == deck


def test_a_colleague_hears_when_a_slide_is_edited(api, two_schools, staff, deck):
    teacher = auth(two_schools["a"]["teacher_token"])
    slide_id = api.get(f"/api/presentations/{deck}", headers=teacher).json()["slides"][0]["id"]

    def edit():
        r = api.patch(
            f"/api/presentations/{deck}/slides/{slide_id}",
            headers=teacher, json={"title": "Edited"},
        )
        assert r.status_code == 200, r.text

    assert _heard(staff["colleague"], edit), "a slide edit reached no colleague"


def test_a_colleague_hears_when_a_deck_leaves_a_dashboard(api, two_schools, staff, deck):
    teacher = auth(two_schools["a"]["teacher_token"])

    def unadopt():
        assert api.delete(f"/api/presentations/{deck}/adopt", headers=teacher).status_code == 204

    assert _heard(staff["colleague"], unadopt), "an un-adoption reached no colleague"


def test_another_school_hears_none_of_it(api, two_schools, staff, deck):
    """Keep next to the tests above: on its own an absence would also pass if
    nothing were delivered to anyone."""
    teacher = auth(two_schools["a"]["teacher_token"])

    def churn():
        assert api.post(f"/api/presentations/{deck}/adopt", headers=teacher).status_code == 201
        _upload(api, two_schools["a"]["teacher_token"], f"{PREFIX} quiet deck")

    leaked = _heard(staff["outsider"], churn, SILENCE_TIMEOUT)
    assert not leaked, f"school B heard about school A's presentations: {leaked}"


def test_a_colleague_hears_when_a_deck_is_deleted(api, two_schools, staff, deck):
    teacher = auth(two_schools["a"]["teacher_token"])

    def delete():
        assert api.delete(f"/api/presentations/{deck}", headers=teacher).status_code == 204

    assert _heard(staff["colleague"], delete), "a deleted deck reached no colleague"
